from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse
from sqlmodel import Session, select
from sqlalchemy import func
from db.database import get_session
from sqlalchemy.exc import IntegrityError
from typing import List

from db.models import Producto, Proveedor, Categoria, Cliente, Ventas, Abono
from db.schemas.clienteSchema import ClienteCreate, ClienteResponse, ClienteUpdate, ClienteBase, AbonoCreate

# --- IMPORTACIONES PARA REPORTES (EXCEL Y PDF SEGURO) ---
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime
from fpdf import FPDF 

router = APIRouter(
    prefix="/clientes",
    tags=["clientes"]
)


@router.get("/", response_model=List[ClienteResponse])
def clientes(session: Session = Depends(get_session)):
    statement = select(Cliente).order_by(Cliente.nombre)
    clientes = session.exec(statement).all()
    return clientes


@router.post("/", response_model=ClienteResponse, status_code=201)
def crear_cliente(cliente: ClienteCreate, session: Session= Depends(get_session)):
    
    # 1. Validar duplicados SOLO si el cliente proporcionó un email
    if cliente.email:
        statement = select(Cliente).where(Cliente.email == cliente.email)
        email_existente = session.exec(statement).first()

        if email_existente:
            raise HTTPException(
                status_code=400,
                detail=f"Error: El email '{cliente.email}' ya está registrado."
            )
    
    # 2. Creación normal
    db_cliente = Cliente.model_validate(cliente)
    try:
        session.add(db_cliente)
        session.commit()
        session.refresh(db_cliente)
    except IntegrityError:
        session.rollback()
        raise HTTPException(
            status_code=400,
            detail=f"Error de integridad: No se pudo registrar al cliente '{db_cliente.nombre}'"
        )

    return db_cliente


@router.patch("/{id}", response_model=ClienteResponse)
def actualizar_cliente(id: int, cliente_data: ClienteUpdate, session: Session = Depends(get_session)):
    db_cliente = session.get(Cliente, id)
    
    if not db_cliente:
        raise HTTPException(status_code=404, detail="Cliente no encontrado")
    datos_nuevos = cliente_data.model_dump(exclude_unset=True)
    
    # Si la propiedad "email" viene en los datos y es un string vacío o puro espacio, se vuelve None.
    if "email" in datos_nuevos and isinstance(datos_nuevos["email"], str):
        if datos_nuevos["email"].strip() == "":
            datos_nuevos["email"] = None
    # -------------------------------------------------------------------------

    db_cliente.sqlmodel_update(datos_nuevos)
    session.add(db_cliente)

    try:
        session.commit()
        session.refresh(db_cliente)
    except IntegrityError:
        session.rollback()
        raise HTTPException(status_code=400, detail="El email ya existe en otro cliente")
    
    return db_cliente


# ===============================================================================
# REGISTRO DE ABONOS (MÉTODO AJUSTADO A TU MODELO REAL 'Abono')
# ===============================================================================
@router.post("/abonos", status_code=201)
def registrar_abono(abono: AbonoCreate, session: Session = Depends(get_session)):
    db_cliente = session.get(Cliente, abono.cliente_id)
    if not db_cliente:
        raise HTTPException(status_code=404, detail="Cliente no encontrado")
    
    if abono.monto_pagado > db_cliente.saldo_deuda:
        raise HTTPException(
            status_code=400, 
            detail=f"El cliente solo debe ${db_cliente.saldo_deuda}. No puede abonar ${abono.monto_pagado}"
        )
        
    db_cliente.saldo_deuda -= abono.monto_pagado
    
    nuevo_registro_abono = Abono(
        cliente_id=abono.cliente_id,
        monto_abonado=abono.monto_pagado, 
        fecha=datetime.now(),
        notas="Abono asentado desde Estado de Cuenta Integral"
    )
    
    session.add(db_cliente)
    session.add(nuevo_registro_abono) 
    session.commit()
    session.refresh(db_cliente)
    
    return {
        "mensaje": "Abono registrado con éxito",
        "cliente": db_cliente.nombre,
        "nuevo_saldo_deuda": db_cliente.saldo_deuda
    }


# ===============================================================================
# CONSULTA HISTÓRICA DE ABONOS 
# ===============================================================================
@router.get("/{id}/abonos")
def obtener_abonos_cliente(id: int, session: Session = Depends(get_session)):
    statement = select(Abono).where(Abono.cliente_id == id).order_by(Abono.fecha.desc())
    return session.exec(statement).all()


# Función auxiliar interna para limpiar acentos en los archivos PDF generados con FPDF
def _limpiar_acentos(texto: str) -> str:
    if not texto:
        return ""
    reemplazos = [("á", "a"), ("é", "e"), ("í", "i"), ("ó", "o"), ("ú", "u"),
                  ("Á", "A"), ("É", "E"), ("Í", "I"), ("Ó", "O"), ("Ú", "U"), ("ñ", "n"), ("Ñ", "N")]
    for original, nuevo in reemplazos:
        texto = texto.replace(original, nuevo)
    return texto


# ===============================================================================
# 1. REPORTE CARTERA INDIVIDUAL: EXCEL
# ===============================================================================
@router.get("/{id}/reporte-excel")
def descargar_excel_cliente(id: int, session: Session = Depends(get_session)):
    db_cliente = session.get(Cliente, id)
    if not db_cliente:
        raise HTTPException(status_code=404, detail="Cliente no registrado")

    statement = select(Ventas).where(Ventas.cliente_id == id).order_by(Ventas.fecha.desc())
    historial_ventas = session.exec(statement).all()

    wb = openpyxl.Workbook()
    ws_summary = wb.active
    ws_summary.title = "Resumen de Cuenta"
    ws_summary.views.sheetView[0].showGridLines = True

    font_title = Font(name="Arial", size=14, bold=True, color="1E293B")
    font_bold = Font(name="Arial", size=11, bold=True, color="1E293B")
    font_regular = Font(name="Arial", size=11, color="334155")
    fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    border_thin = Border(left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
                         top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1'))

    ws_summary['B2'] = "ESTADO DE CUENTA INTEGRAL"
    ws_summary['B2'].font = font_title
    ws_summary['B3'] = f"Reporte Cortado al: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws_summary['B3'].font = Font(name="Arial", size=10, italic=True, color="64748B")

    ws_summary['B5'] = "Nombre del Cliente:"
    ws_summary['B5'].font = font_bold
    ws_summary['C5'] = f"{db_cliente.nombre} {db_cliente.apellido or ''}"
    ws_summary['C5'].font = font_regular

    ws_summary['B6'] = "ID del Sistema:"
    ws_summary['B6'].font = font_bold
    ws_summary['C6'] = db_cliente.id
    ws_summary['C6'].font = font_regular

    ws_summary['B8'] = "SALDO PENDIENTE ACTUAL:"
    ws_summary['B8'].font = font_bold
    ws_summary['C8'] = float(db_cliente.saldo_deuda)
    ws_summary['C8'].font = Font(name="Arial", size=13, bold=True, color="B91C1C")
    ws_summary['C8'].number_format = '$#,##0'
    ws_summary['C8'].fill = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")
    ws_summary['C8'].border = border_thin

    ws_movs = wb.create_sheet(title="Historial Facturas")
    ws_movs.views.sheetView[0].showGridLines = True
    ws_movs['B2'] = "RELACIÓN GENERAL DE FACTURAS EMITIDAS"
    ws_movs['B2'].font = font_title

    headers = ["ID Factura", "Fecha de Emisión", "Tipo de Venta", "Método de Pago", "Monto de Venta"]
    for col_idx, h in enumerate(headers, start=2):
        cell = ws_movs.cell(row=4, column=col_idx, value=h)
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center")

    for idx, v in enumerate(historial_ventas, start=5):
        ws_movs.cell(row=idx, column=2, value=f"#{v.id}").alignment = Alignment(horizontal="center")
        ws_movs.cell(row=idx, column=3, value=v.fecha.strftime('%d/%m/%Y')).alignment = Alignment(horizontal="center")
        ws_movs.cell(row=idx, column=4, value="A Crédito" if v.es_credito else "Contado").alignment = Alignment(horizontal="center")
        ws_movs.cell(row=idx, column=5, value=v.metodo_pago).alignment = Alignment(horizontal="center")
        
        celda_total = ws_movs.cell(row=idx, column=6, value=float(v.total_venta))
        celda_total.number_format = '$#,##0'
        celda_total.alignment = Alignment(horizontal="right")
        celda_total.font = font_regular
        
        for col in range(2, 7):
            ws_movs.cell(row=idx, column=col).border = border_thin

    for sheet in [ws_summary, ws_movs]:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=Establishment_Cuenta_{db_cliente.nombre}.xlsx"}
    )


# ===============================================================================
# 2. REPORTE CARTERA INDIVIDUAL: PDF 
# ===============================================================================
@router.get("/{id}/reporte-pdf")
def descargar_pdf_cliente(id: int, session: Session = Depends(get_session)):
    db_cliente = session.get(Cliente, id)
    if not db_cliente:
        raise HTTPException(status_code=404, detail="Cliente no registrado")

    statement = select(Ventas).where(Ventas.cliente_id == id).order_by(Ventas.fecha.desc())
    historial_ventas = session.exec(statement).all()

    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.add_page()
    pdf.set_margins(15, 20, 15)
    
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, "ESTADO DE CUENTA INTEGRAL", align="L")
    pdf.ln(10) 
    
    pdf.set_font("Arial", "I", 10)
    pdf.set_text_color(100, 116, 139)
    pdf.cell(0, 5, "Historial de Facturacion y Saldos en Cartera - INVENTARIO PRO", align="L")
    pdf.ln(13)
    
    pdf.set_text_color(30, 41, 59)
    pdf.set_fill_color(248, 250, 252)
    pdf.set_font("Arial", "B", 11)
    pdf.cell(0, 8, " INFORMACION DEL CLIENTE", fill=True)
    pdf.ln(8)
    
    pdf.set_font("Arial", "", 10)
    nombre_limpio = _limpiar_acentos(f"{db_cliente.nombre} {db_cliente.apellido or ''}")
    
    pdf.cell(40, 7, "Cliente:", border="B")
    pdf.cell(55, 7, nombre_limpio, border="B")
    pdf.cell(40, 7, "ID Unico:", border="B")
    pdf.cell(0, 7, f"#{db_cliente.id}", border="B")
    pdf.ln(7)
    
    pdf.cell(40, 7, "Telefono:", border="B")
    pdf.cell(55, 7, f"{db_cliente.telefono or 'No registrado'}", border="B")
    pdf.cell(40, 7, "SALDO EN DEUDA:", border="B")
    
    pdf.set_text_color(185, 28, 28) 
    pdf.set_font("Arial", "B", 10)
    pdf.cell(0, 7, f" ${float(db_cliente.saldo_deuda):,.0f} COP", border="B")
    pdf.ln(15)
    
    pdf.set_text_color(255, 255, 255)
    pdf.set_fill_color(30, 41, 59) 
    pdf.set_font("Arial", "B", 9)
    
    pdf.cell(25, 9, "FACTURA", border=1, align="C", fill=True)
    pdf.cell(45, 9, "FECHA EMISION", border=1, align="C", fill=True)
    pdf.cell(35, 9, "CONDICION", border=1, align="C", fill=True)
    pdf.cell(40, 9, "METODO PAGO", border=1, align="C", fill=True)
    pdf.cell(35, 9, "TOTAL BRUTO", border=1, align="C", fill=True)
    pdf.ln(9)
    
    pdf.set_text_color(51, 65, 85)
    pdf.set_font("Arial", "", 10)
    
    total_acumulado = 0.0
    for idx, v in enumerate(historial_ventas):
        total_acumulado += float(v.total_venta)
        fill_bg = True if idx % 2 == 0 else False
        pdf.set_fill_color(248, 250, 252)
        
        metodo_limpio = _limpiar_acentos(v.metodo_pago)
        condicion = "A Credito" if v.es_credito else "Contado"
        
        pdf.cell(25, 8, f"#{v.id}", border=1, align="C", fill=fill_bg)
        pdf.cell(45, 8, v.fecha.strftime('%d/%m/%Y'), border=1, align="C", fill=fill_bg)
        pdf.cell(35, 8, condicion, border=1, align="C", fill=fill_bg)
        pdf.cell(40, 8, metodo_limpio, border=1, align="C", fill=fill_bg)
        pdf.cell(35, 8, f"${float(v.total_venta):,.0f}", border=1, align="R", fill=fill_bg)
        pdf.ln(8)

    pdf.set_font("Arial", "B", 10)
    pdf.set_fill_color(226, 232, 240)
    pdf.cell(145, 9, "TOTAL HISTORICO FACTURADO  ", border=1, align="R", fill=True)
    pdf.cell(35, 9, f"${total_acumulado:,.0f}", border=1, align="R", fill=True)
    pdf.ln(9)

    pdf_bytes = pdf.output()
    stream = BytesIO(pdf_bytes)

    return StreamingResponse(
        stream,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=Estado_Cuenta_{db_cliente.nombre}.pdf"}
    )


# ===============================================================================
# 3. REPORTE INVENTARIO GLOBAL - EXCEL (CORREGIDO DE RAÍZ A PRECIO DE COMPRA)
# ===============================================================================
@router.get("/global/inventario-excel")
def excel_valorizacion_inventario(session: Session = Depends(get_session)):
    productos = session.exec(select(Producto)).all()
    categorias = session.exec(select(Categoria)).all()
    mapa_cat = {c.id: c.nombre for c in categorias}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Valorizacion de Stock"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Arial", size=14, bold=True, color="1E293B")
    font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Arial", size=11, bold=True, color="1E293B")
    font_regular = Font(name="Arial", size=11, color="334155")
    fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_total = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    border_thin = Border(left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
                         top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1'))

    ws['B2'] = "REPORTE GLOBAL DE VALORIZACIÓN DE INVENTARIO"
    ws['B2'].font = font_title
    ws['B3'] = f"Corte realizado el: {datetime.now().strftime('%d/%m/%Y %I:%M %p')}"
    ws['B3'].font = Font(name="Arial", size=10, italic=True, color="64748B")

    # MODIFICADO: Ajustados los encabezados a Costo Real e Inversión
    headers = ["ID Ref", "Nombre del Artículo", "Categoría", "Stock Físico", "Costo Unit. (Compra)", "Inversión Total"]
    for col_idx, text in enumerate(headers, start=2):
        cell = ws.cell(row=5, column=col_idx, value=text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    total_unidades = 0
    total_capital_estimado = 0.0

    for idx, p in enumerate(productos, start=6):
        # MODIFICADO: Multiplica por precio_compra (Costo de adquisición)
        valor_total_item = p.stock * float(p.precio_compra)
        total_unidades += p.stock
        total_capital_estimado += valor_total_item

        ws.cell(row=idx, column=2, value=f"#{p.id}").alignment = Alignment(horizontal="center")
        ws.cell(row=idx, column=3, value=p.nombre).font = font_regular
        ws.cell(row=idx, column=4, value=mapa_cat.get(p.categoria_id, "Sin Categoría")).font = font_regular
        
        c_stock = ws.cell(row=idx, column=5, value=p.stock)
        c_stock.number_format = '#,##0" unid."'
        c_stock.alignment = Alignment(horizontal="center")

        # MODIFICADO: Mapea el precio_compra en la celda
        c_precio = ws.cell(row=idx, column=6, value=float(p.precio_compra))
        c_precio.number_format = '$#,##0'
        c_precio.alignment = Alignment(horizontal="right")

        c_val = ws.cell(row=idx, column=7, value=valor_total_item)
        c_val.number_format = '$#,##0'
        c_val.alignment = Alignment(horizontal="right")
        c_val.font = font_bold

        for col in range(2, 8):
            ws.cell(row=idx, column=col).border = border_thin

    tot_row = len(productos) + 6
    ws.cell(row=tot_row, column=2, value="TOTAL CONSOLIDADO").font = font_bold
    ws.cell(row=tot_row, column=5, value=total_unidades).font = font_bold
    ws.cell(row=tot_row, column=5).number_format = '#,##0" unid."'
    ws.cell(row=tot_row, column=7, value=total_capital_estimado).font = Font(name="Arial", size=11, bold=True, color="1E3A8A")
    ws.cell(row=tot_row, column=7).number_format = '$#,##0'

    for col in range(2, 8):
        c = ws.cell(row=tot_row, column=col)
        c.fill = fill_total
        c.border = Border(top=Side(style='thin', color='1E293B'), bottom=Side(style='double', color='1E293B'))

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=Valorizacion_Inventario_Global.xlsx"}
    )


# ===============================================================================
# 4. REPORTE INVENTARIO GLOBAL - PDF (CORREGIDO DE RAÍZ A PRECIO DE COMPRA)
# ===============================================================================
@router.get("/global/inventario-pdf")
def pdf_valorizacion_inventario(session: Session = Depends(get_session)):
    productos = session.exec(select(Producto)).all()
    categorias = session.exec(select(Categoria)).all()
    mapa_cat = {c.id: c.nombre for c in categorias}

    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.add_page()
    pdf.set_margins(15, 20, 15)
    
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, "VALORIZACION DE INVENTARIO CENTRAL", align="L")
    pdf.ln(10)
    
    pdf.set_font("Arial", "I", 10)
    pdf.set_text_color(100, 116, 139)
    pdf.cell(0, 5, f"Reporte de Activos Fisicos - Generado: {datetime.now().strftime('%d/%m/%Y')}", align="L")
    pdf.ln(15)
    
    pdf.set_text_color(255, 255, 255)
    pdf.set_fill_color(30, 41, 59)
    pdf.set_font("Arial", "B", 9)
    
    pdf.cell(20, 9, "REF", border=1, align="C", fill=True)
    pdf.cell(55, 9, "DESCRIPCION", border=1, align="L", fill=True)
    pdf.cell(30, 9, "CATEGORIA", border=1, align="C", fill=True)
    pdf.cell(25, 9, "STOCK", border=1, align="C", fill=True)
    # MODIFICADO: Títulos alineados a Costo Real de Inversión
    pdf.cell(25, 9, "COSTO UNIT.", border=1, align="C", fill=True)
    pdf.cell(25, 9, "TOTAL COSTO", border=1, align="C", fill=True)
    pdf.ln(9)
    
    pdf.set_text_color(51, 65, 85)
    pdf.set_font("Arial", "", 9)
    
    total_unidades = 0
    total_valor = 0.0
    
    for idx, p in enumerate(productos):
        # MODIFICADO: Cálculo estructurado sobre el precio_compra
        subtotal = p.stock * float(p.precio_compra)
        total_unidades += p.stock
        total_valor += subtotal
        
        fill_bg = True if idx % 2 == 0 else False
        pdf.set_fill_color(248, 250, 252)
        
        nombre_p = _limpiar_acentos(p.nombre)
        cat_p = _limpiar_acentos(mapa_cat.get(p.categoria_id, "General"))
        
        pdf.cell(20, 8, f"#{p.id}", border=1, align="C", fill=fill_bg)
        pdf.cell(55, 8, nombre_p, border=1, align="L", fill=fill_bg)
        pdf.cell(30, 8, cat_p, border=1, align="C", fill=fill_bg)
        pdf.cell(25, 8, f"{p.stock} u.", border=1, align="C", fill=fill_bg)
        # MODIFICADO: Inyecta el precio_compra en el renglón
        pdf.cell(25, 8, f"${float(p.precio_compra):,.0f}", border=1, align="R", fill=fill_bg)
        pdf.cell(25, 8, f"${subtotal:,.0f}", border=1, align="R", fill=fill_bg)
        pdf.ln(8)
        
    pdf.set_font("Arial", "B", 10)
    pdf.set_fill_color(226, 232, 240)
    pdf.cell(105, 9, "TOTALES CONSOLIDADOS  ", border=1, align="R", fill=True)
    pdf.cell(25, 9, f"{total_unidades} u.", border=1, align="C", fill=True)
    pdf.cell(25, 9, "", border=1, fill=True)
    pdf.cell(25, 9, f"${total_valor:,.0f}", border=1, align="R", fill=True)
    pdf.ln(9)
    
    pdf_bytes = pdf.output()
    return StreamingResponse(BytesIO(pdf_bytes), media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=Valorizacion_Inventario_Global.pdf"})


# ===============================================================================
# 5. REPORTE DE VENTAS POR RANGO DE FECHAS - EXCEL
# ===============================================================================
@router.get("/global/ventas-excel")
def excel_ventas_por_fechas(fecha_inicio: str, fecha_fin: str, session: Session = Depends(get_session)):
    try:
        f_inicio = datetime.strptime(fecha_inicio, "%Y-%m-%d")
        f_fin = datetime.strptime(fecha_fin, "%Y-%m-%d").replace(hour=23, minute=59, second=59)
    except ValueError:
        raise HTTPException(status_code=400, detail="Formato de fecha invalido. Use YYYY-MM-DD")

    statement = select(Ventas).where(Ventas.fecha >= f_inicio, Ventas.fecha <= f_fin).order_by(Ventas.fecha.asc())
    historial_ventas = session.exec(statement).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Balance de Ventas"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Arial", size=14, bold=True, color="1E293B")
    font_bold = Font(name="Arial", size=11, bold=True, color="1E293B")
    font_regular = Font(name="Arial", size=11, color="334155")
    fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_total = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    border_thin = Border(left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
                         top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1'))

    ws['B2'] = "REPORTE CONSOLIDADO DE VENTAS"
    ws['B2'].font = font_title
    ws['B3'] = f"Periodo auditado: Del {f_inicio.strftime('%d/%m/%Y')} al {f_fin.strftime('%d/%m/%Y')}"
    ws['B3'].font = Font(name="Arial", size=10, italic=True, color="64748B")

    headers = ["ID Venta", "Fecha y Hora", "Condicion", "Metodo Pago", "Total Facturado"]
    for col_idx, text in enumerate(headers, start=2):
        cell = ws.cell(row=5, column=col_idx, value=text)
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center")

    total_dinero_ventas = 0.0

    for idx, v in enumerate(historial_ventas, start=6):
        total_dinero_ventas += float(v.total_venta)

        ws.cell(row=idx, column=2, value=f"#{v.id}").alignment = Alignment(horizontal="center")
        ws.cell(row=idx, column=3, value=v.fecha.strftime('%d/%m/%Y %I:%M %p')).alignment = Alignment(horizontal="center")
        ws.cell(row=idx, column=4, value="A Credito" if v.es_credito else "Contado").alignment = Alignment(horizontal="center")
        ws.cell(row=idx, column=5, value=v.metodo_pago).alignment = Alignment(horizontal="center")
        
        c_total = ws.cell(row=idx, column=6, value=float(v.total_venta))
        c_total.number_format = '$#,##0'
        c_total.alignment = Alignment(horizontal="right")
        c_total.font = font_regular

        for col in range(2, 7):
            ws.cell(row=idx, column=col).border = border_thin

    tot_row = len(historial_ventas) + 6
    ws.cell(row=tot_row, column=2, value="TOTAL INGRESOS BRUTOS").font = font_bold
    ws.cell(row=tot_row, column=6, value=total_dinero_ventas).font = Font(name="Arial", size=11, bold=True, color="16A34A")
    ws.cell(row=tot_row, column=6).number_format = '$#,##0'

    for col in range(2, 7):
        c = ws.cell(row=tot_row, column=col)
        c.fill = fill_total
        c.border = Border(top=Side(style='thin', color='1E293B'), bottom=Side(style='double', color='1E293B'))

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(stream, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=Reporte_Ventas.xlsx"})


# ===============================================================================
# 6. REPORTE DE VENTAS POR RANGO DE FECHAS - PDF 
# ===============================================================================
@router.get("/global/ventas-pdf")
def pdf_ventas_por_fechas(fecha_inicio: str, fecha_fin: str, session: Session = Depends(get_session)):
    try:
        f_inicio = datetime.strptime(fecha_inicio, "%Y-%m-%d")
        f_fin = datetime.strptime(fecha_fin, "%Y-%m-%d").replace(hour=23, minute=59, second=59)
    except ValueError:
        raise HTTPException(status_code=400, detail="Formato de fecha invalido. Use YYYY-MM-DD")

    statement = select(Ventas).where(Ventas.fecha >= f_inicio, Ventas.fecha <= f_fin).order_by(Ventas.fecha.asc())
    historial_ventas = session.exec(statement).all()

    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.add_page()
    pdf.set_margins(15, 20, 15)
    
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, "REPORTE AUDITABLE DE VENTAS", align="L")
    pdf.ln(10)
    
    pdf.set_font("Arial", "I", 10)
    pdf.set_text_color(100, 116, 139)
    pdf.cell(0, 5, f"Periodo: {f_inicio.strftime('%d/%m/%Y')} al {f_fin.strftime('%d/%m/%Y')} - INVENTARIO PRO", align="L")
    pdf.ln(15)
    
    pdf.set_text_color(255, 255, 255)
    pdf.set_fill_color(30, 41, 59)
    pdf.set_font("Arial", "B", 9)
    
    pdf.cell(20, 9, "ID VENTA", border=1, align="C", fill=True)
    pdf.cell(50, 9, "FECHA Y HORA", border=1, align="C", fill=True)
    pdf.cell(35, 9, "CONDICION", border=1, align="C", fill=True)
    pdf.cell(40, 9, "METODO PAGO", border=1, align="C", fill=True)
    pdf.cell(35, 9, "VALOR FACTURADO", border=1, align="C", fill=True)
    pdf.ln(9)
    
    pdf.set_text_color(51, 65, 85)
    pdf.set_font("Arial", "", 9)
    
    total_dinero = 0.0
    for idx, v in enumerate(historial_ventas):
        total_dinero += float(v.total_venta)
        fill_bg = True if idx % 2 == 0 else False
        pdf.set_fill_color(248, 250, 252)
        
        metodo = _limpiar_acentos(v.metodo_pago)
        condicion = "A Credito" if v.es_credito else "Contado"
        
        pdf.cell(20, 8, f"#{v.id}", border=1, align="C", fill=fill_bg)
        pdf.cell(50, 8, v.fecha.strftime('%d/%m/%Y %I:%M %p'), border=1, align="C", fill=fill_bg)
        pdf.cell(35, 8, condicion, border=1, align="C", fill=fill_bg)
        pdf.cell(40, 8, metodo, border=1, align="C", fill=fill_bg)
        pdf.cell(35, 8, f"${float(v.total_venta):,.0f}", border=1, align="R", fill=fill_bg)
        pdf.ln(8)
        
    pdf.set_font("Arial", "B", 10)
    pdf.set_fill_color(226, 232, 240)
    pdf.cell(146, 9, "TOTAL INGRESOS BRUTOS  ", border=1, align="R", fill=True)
    pdf.cell(35, 9, f"${total_dinero:,.0f}", border=1, align="R", fill=True)
    pdf.ln(9)
    
    pdf_bytes = pdf.output()
    return StreamingResponse(BytesIO(pdf_bytes), media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=Reporte_Ventas_Global.pdf"})


# ===============================================================================
# 7. REPORTE CONSOLIDADO DE DEUDORES (CARTERA) - EXCEL
# ===============================================================================
@router.get("/global/deudores-excel")
def excel_consolidado_deudores(session: Session = Depends(get_session)):
    statement = select(Cliente).where(Cliente.saldo_deuda > 0).order_by(Cliente.saldo_deuda.desc())
    deudores = session.exec(statement).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cartera por Cobrar"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Arial", size=14, bold=True, color="1E293B")
    font_bold = Font(name="Arial", size=11, bold=True, color="1E293B")
    font_regular = Font(name="Arial", size=11, color="334155")
    fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_total = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")
    border_thin = Border(left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
                         top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1'))

    ws['B2'] = "REPORTE CONSOLIDADO DE CARTERA (DEUDORES)"
    ws['B2'].font = font_title
    ws['B3'] = f"Reporte generado el: {datetime.now().strftime('%d/%m/%Y %I:%M %p')} - Filtro: Saldos activos"
    ws['B3'].font = Font(name="Arial", size=10, italic=True, color="64748B")

    headers = ["ID", "Nombre Completo", "Celular / Contacto", "Correo Electrónico", "Saldo en Mora"]
    for col_idx, text in enumerate(headers, start=2):
        cell = ws.cell(row=5, column=col_idx, value=text)
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center")

    total_cartera_recuperar = 0.0

    for idx, c in enumerate(deudores, start=6):
        saldo_val = float(c.saldo_deuda)
        total_cartera_recuperar += saldo_val

        ws.cell(row=idx, column=2, value=f"#{c.id}").alignment = Alignment(horizontal="center")
        ws.cell(row=idx, column=3, value=f"{c.nombre} {c.apellido or ''}").font = font_regular
        ws.cell(row=idx, column=4, value=c.telefono or "Sin número").alignment = Alignment(horizontal="center")
        ws.cell(row=idx, column=5, value=c.email).font = font_regular
        
        c_saldo = ws.cell(row=idx, column=6, value=saldo_val)
        c_saldo.number_format = '$#,##0'
        c_saldo.alignment = Alignment(horizontal="right")
        c_saldo.font = Font(name="Arial", size=11, bold=True, color="991B1B")

        for col in range(2, 7):
            ws.cell(row=idx, column=col).border = border_thin

    tot_row = len(deudores) + 6
    ws.cell(row=tot_row, column=2, value="TOTAL CARTERA POR RECUPERAR").font = font_bold
    ws.cell(row=tot_row, column=6, value=total_cartera_recuperar).font = Font(name="Arial", size=11, bold=True, color="B91C1C")
    ws.cell(row=tot_row, column=6).number_format = '$#,##0'

    for col in range(2, 7):
        cell_t = ws.cell(row=tot_row, column=col)
        cell_t.fill = fill_total
        cell_t.border = Border(top=Side(style='thin', color='B91C1C'), bottom=Side(style='double', color='B91C1C'))

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(stream, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=Consolidado_Deudores.xlsx"})


# ===============================================================================
# 8. REPORTE CONSOLIDADO DE DEUDORES (CARTERA) - PDF
# ===============================================================================
@router.get("/global/deudores-pdf")
def pdf_consolidado_deudores(session: Session = Depends(get_session)):
    statement = select(Cliente).where(Cliente.saldo_deuda > 0).order_by(Cliente.saldo_deuda.desc())
    deudores = session.exec(statement).all()

    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.add_page()
    pdf.set_margins(15, 20, 15)
    
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, "REPORTE CONSOLIDADO DE DEUDORES", align="L")
    pdf.ln(10)
    
    pdf.set_font("Arial", "I", 10)
    pdf.set_text_color(100, 116, 139)
    pdf.cell(0, 5, f"Control de Saldos en Mora y Cuentas por Cobrar - Corte: {datetime.now().strftime('%d/%m/%Y')}", align="L")
    pdf.ln(15)
    
    pdf.set_text_color(255, 255, 255)
    pdf.set_fill_color(30, 41, 59)
    pdf.set_font("Arial", "B", 9)
    
    pdf.cell(15, 9, "ID", border=1, align="C", fill=True)
    pdf.cell(55, 9, "NOMBRE DEL CLIENTE", border=1, align="L", fill=True)
    pdf.cell(30, 9, "TELEFONO", border=1, align="C", fill=True)
    pdf.cell(50, 9, "CORREO ELECTRONICO", border=1, align="L", fill=True)
    pdf.cell(30, 9, "SALDO MORA", border=1, align="C", fill=True)
    pdf.ln(9)
    
    pdf.set_text_color(51, 65, 85)
    pdf.set_font("Arial", "", 9)
    
    total_cartera = 0.0
    for idx, c in enumerate(deudores):
        saldo_val = float(c.saldo_deuda)
        total_cartera += saldo_val
        
        fill_bg = True if idx % 2 == 0 else False
        pdf.set_fill_color(248, 250, 252)
        
        nombre_c = _limpiar_acentos(f"{c.nombre} {c.apellido or ''}")
        tel_c = c.telefono or "Sin numero"
        
        pdf.cell(15, 8, f"#{c.id}", border=1, align="C", fill=fill_bg)
        pdf.cell(55, 8, nombre_c, border=1, align="L", fill=fill_bg)
        pdf.cell(30, 8, tel_c, border=1, align="C", fill=fill_bg)
        pdf.cell(50, 8, c.email, border=1, align="L", fill=fill_bg)
        
        pdf.set_text_color(185, 28, 28)
        pdf.set_font("Arial", "B", 9)
        pdf.cell(30, 8, f"${saldo_val:,.0f}", border=1, align="R", fill=fill_bg)
        
        pdf.set_text_color(51, 65, 85)
        pdf.set_font("Arial", "", 9)
        pdf.ln(8)
        
    pdf.set_font("Arial", "B", 10)
    pdf.set_fill_color(254, 242, 242) 
    pdf.set_text_color(153, 27, 27) 
    pdf.cell(150, 9, "TOTAL GENERAL DE CARTERA POR RECUPERAR  ", border=1, align="R", fill=True)
    pdf.cell(30, 9, f"${total_cartera:,.0f}", border=1, align="R", fill=True)
    pdf.ln(9)
    
    pdf_bytes = pdf.output()
    return StreamingResponse(BytesIO(pdf_bytes), media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=Consolidado_Deudores.pdf"})